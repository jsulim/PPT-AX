from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

from core.pptx.slots import extract_text_slots


def test_extracts_text_box_placeholder_and_table_slots(tmp_path: Path) -> None:
    pptx_path = tmp_path / "template.pptx"
    presentation = Presentation()

    slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    slide.shapes.title.text = "기존 사업명"
    slide.placeholders[1].text = "기존 부제"

    second_slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    textbox = second_slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    textbox.text_frame.text = "추진전략 원문"

    table_shape = second_slide.shapes.add_table(1, 2, Inches(1), Inches(2), Inches(5), Inches(1))
    table = table_shape.table
    table.cell(0, 0).text = "기관명"
    table.cell(0, 1).text = "기존 기관"

    footer = second_slide.shapes.add_textbox(Inches(1), Inches(6.9), Inches(1), Inches(0.3))
    footer.text_frame.text = "1"

    presentation.save(pptx_path)

    result = extract_text_slots(str(pptx_path))

    texts = {slot.original_text for slot in result.slots}
    assert result.slide_count == 2
    assert "기존 사업명" in texts
    assert "기존 부제" in texts
    assert "추진전략 원문" in texts
    assert "기관명" in texts
    assert "기존 기관" in texts
    assert any(slot.kind == "table_cell" for slot in result.slots)
    assert any(slot.original_text == "1" and slot.locked for slot in result.slots)
    assert all(slot.max_char_count >= slot.char_count for slot in result.slots)
