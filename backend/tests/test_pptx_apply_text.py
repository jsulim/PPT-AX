from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

from core.pptx.apply_text import TextReplacement, apply_text_replacements
from core.pptx.slots import extract_text_slots


def test_apply_text_replacements_preserves_original_file(tmp_path: Path) -> None:
    template_path = tmp_path / "template.pptx"
    output_path = tmp_path / "output.pptx"

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    textbox.text_frame.text = "OLD PROJECT"
    table_shape = slide.shapes.add_table(1, 1, Inches(1), Inches(2), Inches(4), Inches(1))
    table_shape.table.cell(0, 0).text = "OLD AGENCY"
    presentation.save(template_path)

    slots = extract_text_slots(str(template_path)).slots
    title_slot = next(slot for slot in slots if slot.original_text == "OLD PROJECT")
    agency_slot = next(slot for slot in slots if slot.original_text == "OLD AGENCY")

    report = apply_text_replacements(
        template_path,
        output_path,
        [
            TextReplacement(title_slot.slide_index, title_slot.shape_id, "NEW PROJECT"),
            TextReplacement(agency_slot.slide_index, agency_slot.shape_id, "NEW AGENCY"),
        ],
    )

    result = Presentation(str(output_path))
    result_text = "\n".join(shape.text for shape in result.slides[0].shapes if shape.has_text_frame)
    result_table_text = result.slides[0].shapes[-1].table.cell(0, 0).text
    original = Presentation(str(template_path))

    assert report.skipped == []
    assert "NEW PROJECT" in result_text
    assert result_table_text == "NEW AGENCY"
    assert "OLD PROJECT" in original.slides[0].shapes[0].text
